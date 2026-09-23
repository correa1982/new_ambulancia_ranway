from flask import Blueprint, render_template, request, redirect, url_for, flash, session, send_file
from db import get_db
from utils import login_required, get_user_info, get_configuracion, ahora
import json
from datetime import datetime, date, time, timedelta
from io import BytesIO
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def register_routes(app):
    @app.route("/formularios/reporte_actividades/nuevo", methods=["GET", "POST"])
    @login_required
    def form_reporte_actividades():
        formularios_acceso = session["usuario"].get("formularios_acceso", {})
        tiene_permiso = session["usuario"].get("rol") == "admin"
        es_admin_vol = session["usuario"].get("rol") == "admin"
        
        if not tiene_permiso:
            if isinstance(formularios_acceso, dict):
                for perfiles_acc in formularios_acceso.values():
                    if 'reporte_actividades' in perfiles_acc:
                        tiene_permiso = True
                    if 'voluntariado_admin' in perfiles_acc:
                        es_admin_vol = True
            elif isinstance(formularios_acceso, list):
                if 'reporte_actividades' in formularios_acceso:
                    tiene_permiso = True
                if 'voluntariado_admin' in formularios_acceso:
                    es_admin_vol = True
                    
        if not tiene_permiso and not es_admin_vol:
            flash("No tienes acceso al formulario de Reporte de Actividades.", "error")
            return redirect(url_for("dashboard"))

        if request.method == "POST":
            nombre_evento = request.form.get("nombre_evento")
            fecha_evento = request.form.get("fecha_evento")
            lugar_evento = request.form.get("lugar_evento")
            hora_inicio = request.form.get("hora_inicio")
            hora_fin = request.form.get("hora_fin")
            tipo_servicio = request.form.get("tipo_servicio")
            
            ambulancia_tab = request.form.get("ambulancia_tab") or "0"
            ambulancia_tam = request.form.get("ambulancia_tam") or "0"
            pasm = request.form.get("pasm") or "0"
            pasb = request.form.get("pasb") or "0"
            equipos_intervencion = request.form.get("equipos_intervencion") or "0"
            moto_aph = request.form.get("moto_aph") or "0"
            unidad_rescate = request.form.get("unidad_rescate") or "0"
            unidad_logistica = request.form.get("unidad_logistica") or "0"
            
            total_personal = request.form.get("total_personal")
            pacientes_atendidos = request.form.get("pacientes_atendidos")
            pacientes_trasladados = request.form.get("pacientes_trasladados")
            observaciones = request.form.get("observaciones")
            
            conn = get_db()
            identificacion = session["usuario"]["identificacion"]
            firma, perfil, _ = get_user_info(conn, identificacion)
            now_str = ahora().strftime("%Y-%m-%d %H:%M:%S")
            registrado_por = session["usuario"]["nombre"]
            
            conn.execute("""
                INSERT INTO reporte_actividades (
                    nombre_evento, fecha_evento, lugar_evento, hora_inicio, hora_fin, tipo_servicio,
                    ambulancia_tab, ambulancia_tam, pasm, pasb, equipos_intervencion, moto_aph, unidad_rescate, unidad_logistica,
                    total_personal, pacientes_atendidos, pacientes_trasladados, observaciones,
                    registrado_por, registrado_por_identificacion, perfil_registrador, firma_registrador, fecha_registro
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                nombre_evento, fecha_evento, lugar_evento, hora_inicio, hora_fin, tipo_servicio,
                ambulancia_tab, ambulancia_tam, pasm, pasb, equipos_intervencion, moto_aph, unidad_rescate, unidad_logistica,
                total_personal, pacientes_atendidos, pacientes_trasladados, observaciones,
                registrado_por, identificacion, perfil, firma, now_str
            ))
            
            conn.commit()
            conn.close()
            
            flash("Reporte de actividades registrado exitosamente.", "success")
            return redirect(url_for("dashboard"))
            
        return render_template("form_reporte_actividades.html", es_admin_vol=es_admin_vol)

    @app.route("/formularios/reporte_actividades/registros")
    @login_required
    def registros_reporte_actividades():
        formularios_acceso = session["usuario"].get("formularios_acceso", {})
        es_admin_vol = session["usuario"].get("rol") == "admin"
        
        if not es_admin_vol:
            if isinstance(formularios_acceso, dict):
                for perfiles_acc in formularios_acceso.values():
                    if 'voluntariado_admin' in perfiles_acc:
                        es_admin_vol = True
                        break
            elif isinstance(formularios_acceso, list):
                if 'voluntariado_admin' in formularios_acceso:
                    es_admin_vol = True

        if not es_admin_vol:
            flash("Acceso denegado. Solo el Admin Voluntariado puede ver estos registros.", "error")
            return redirect(url_for("dashboard"))

        conn = get_db()
        registros = conn.execute("SELECT * FROM reporte_actividades WHERE estado IN ('Pendiente', 'Avalado') ORDER BY id DESC").fetchall()
        conn.close()
        
        return render_template("registros_reporte_actividades.html", registros=registros)
        
    @app.route("/formularios/reporte_actividades/anulados")
    @app.route("/formularios/reporte_actividades/inactivos")
    @login_required
    def anulados_reporte_actividades():
        formularios_acceso = session["usuario"].get("formularios_acceso", {})
        es_admin_vol = session["usuario"].get("rol") == "admin"
        
        if not es_admin_vol:
            if isinstance(formularios_acceso, dict):
                for perfiles_acc in formularios_acceso.values():
                    if 'voluntariado_admin' in perfiles_acc:
                        es_admin_vol = True
                        break
            elif isinstance(formularios_acceso, list):
                if 'voluntariado_admin' in formularios_acceso:
                    es_admin_vol = True

        if not es_admin_vol:
            flash("Acceso denegado.", "error")
            return redirect(url_for("dashboard"))

        conn = get_db()
        registros = conn.execute("SELECT * FROM reporte_actividades WHERE estado IN ('Anulado', 'Inactivo') ORDER BY id DESC").fetchall()
        conn.close()
        
        return render_template("anulados_reporte_actividades.html", registros=registros)

    inactivos_reporte_actividades = anulados_reporte_actividades

    @app.route("/formularios/reporte_actividades/avalar/<int:record_id>", methods=["POST"])
    @login_required
    def avalar_reporte(record_id):
        es_admin_vol = session["usuario"].get("rol") == "admin"
        formularios_acceso = session["usuario"].get("formularios_acceso", {})
        if not es_admin_vol:
            if isinstance(formularios_acceso, dict):
                for perfiles_acc in formularios_acceso.values():
                    if 'voluntariado_admin' in perfiles_acc:
                        es_admin_vol = True
                        break
            elif isinstance(formularios_acceso, list):
                if 'voluntariado_admin' in formularios_acceso:
                    es_admin_vol = True

        if not es_admin_vol:
            flash("No tienes permiso.", "error")
            return redirect(url_for("dashboard"))

        conn = get_db()
        firma, _, _ = get_user_info(conn, session["usuario"]["identificacion"])
        now_str = ahora().strftime("%Y-%m-%d %H:%M:%S")
        
        avalado_por = session["usuario"]["nombre"]
        avalado_por_identificacion = session["usuario"]["identificacion"]
        
        conn.execute("""
            UPDATE reporte_actividades 
            SET estado = 'Avalado', avalado_por = ?, avalado_por_identificacion = ?, firma_avalador = ?, fecha_aval = ?
            WHERE id = ?
        """, (avalado_por, avalado_por_identificacion, firma, now_str, record_id))
        
        conn.commit()
        conn.close()
        
        flash("Reporte avalado exitosamente.", "success")
        return redirect(url_for("registros_reporte_actividades"))
        
    @app.route("/formularios/reporte_actividades/anular/<int:record_id>", methods=["POST"])
    @app.route("/formularios/reporte_actividades/inactivar/<int:record_id>", methods=["POST"])
    @login_required
    def anular_reporte(record_id):
        es_admin_vol = session["usuario"].get("rol") == "admin"
        formularios_acceso = session["usuario"].get("formularios_acceso", {})
        if not es_admin_vol:
            if isinstance(formularios_acceso, dict):
                for perfiles_acc in formularios_acceso.values():
                    if 'voluntariado_admin' in perfiles_acc:
                        es_admin_vol = True
                        break
            elif isinstance(formularios_acceso, list):
                if 'voluntariado_admin' in formularios_acceso:
                    es_admin_vol = True

        if not es_admin_vol:
            flash("No tienes permiso.", "error")
            return redirect(url_for("dashboard"))

        conn = get_db()
        conn.execute("UPDATE reporte_actividades SET estado = 'Anulado' WHERE id = ?", (record_id,))
        conn.commit()
        conn.close()
        
        flash("Reporte anulado correctamente. No afectará las estadísticas.", "warning")
        return redirect(url_for("registros_reporte_actividades"))

    inactivar_reporte = anular_reporte

    @app.route("/formularios/reporte_actividades/imprimir/<int:record_id>")
    @login_required
    def imprimir_reporte(record_id):
        conn = get_db()
        registro = conn.execute("SELECT * FROM reporte_actividades WHERE id = ?", (record_id,)).fetchone()
        
        if not registro:
            conn.close()
            flash("Registro no encontrado.", "error")
            return redirect(url_for("dashboard"))
            
        cfg = get_configuracion(conn)
        conn.close()

        # ── Formatear horas como "HH:MM" independientemente del tipo del driver ──
        import datetime as _dt

        def _fmt_hora(val):
            if val is None:
                return ""
            if isinstance(val, _dt.time):
                return val.strftime("%H:%M")
            if isinstance(val, _dt.timedelta):
                total_sec = int(val.total_seconds())
                h = (total_sec // 3600) % 24
                m = (total_sec % 3600) // 60
                return f"{h:02d}:{m:02d}"
            s = str(val).strip()
            # Si viene como "HH:MM:SS" recortar a "HH:MM"
            if len(s) >= 5:
                return s[:5]
            return s

        # Convertir la Row a dict mutable para poder agregar campos formateados
        registro_dict = dict(registro)
        registro_dict["hora_inicio_fmt"] = _fmt_hora(registro_dict.get("hora_inicio"))
        registro_dict["hora_fin_fmt"]    = _fmt_hora(registro_dict.get("hora_fin"))

        return render_template("imprimir_reporte_actividades.html",
                               registro=registro_dict, cfg=cfg)

    @app.route("/formularios/reporte_actividades/estadisticas")
    @login_required
    def estadisticas_reporte():
        es_admin_vol = session["usuario"].get("rol") == "admin"
        formularios_acceso = session["usuario"].get("formularios_acceso", {})
        
        if not es_admin_vol:
            if isinstance(formularios_acceso, dict):
                for perfiles_acc in formularios_acceso.values():
                    if 'voluntariado_admin' in perfiles_acc:
                        es_admin_vol = True
                        break
            elif isinstance(formularios_acceso, list):
                if 'voluntariado_admin' in formularios_acceso:
                    es_admin_vol = True

        if not es_admin_vol:
            flash("No tienes permisos.", "error")
            return redirect(url_for("dashboard"))

        conn = get_db()
        todos      = conn.execute("SELECT * FROM reporte_actividades").fetchall()
        conn.close()

        avalados   = [r for r in todos if r.get("estado") == "Avalado"]
        anulados   = [r for r in todos if r.get("estado") in ("Anulado", "Inactivo")]
        pendientes = [r for r in todos if r.get("estado") not in ("Avalado", "Anulado", "Inactivo")]

        # ── Resumen de estados ──
        resumen_estados = {
            "total":      len(todos),
            "avalados":   len(avalados),
            "anulados":   len(anulados),
            "pendientes": len(pendientes),
        }

        # ── Estadísticas por tipo de servicio (solo Avalados) ──
        import datetime as _dt

        def _to_minutes(val):
            """Convierte hora_inicio / hora_fin al total de minutos desde medianoche,
            independientemente de si el driver devuelve str, datetime.time o timedelta."""
            if isinstance(val, _dt.time):
                return val.hour * 60 + val.minute
            if isinstance(val, _dt.timedelta):
                total = int(val.total_seconds())
                return (total // 60) % (24 * 60)
            if isinstance(val, str):
                parts = val.strip().split(":")
                return int(parts[0]) * 60 + int(parts[1])
            return None

        RECURSOS = ["ambulancia_tab","ambulancia_tam","pasm","pasb",
                    "equipos_intervencion","moto_aph","unidad_rescate","unidad_logistica"]
        stats_servicios = {}
        for r in avalados:
            tipo = r.get("tipo_servicio")
            if not tipo:
                continue
            if tipo not in stats_servicios:
                stats_servicios[tipo] = {
                    "cantidad": 0, "horas": 0.0,
                    "total_personal": 0, "pacientes_atendidos": 0,
                    "pacientes_trasladados": 0,
                    **{k: 0 for k in RECURSOS}
                }
            s = stats_servicios[tipo]
            s["cantidad"]              += 1
            s["total_personal"]        += int(r["total_personal"]        or 0)
            s["pacientes_atendidos"]   += int(r["pacientes_atendidos"]   or 0)
            s["pacientes_trasladados"] += int(r["pacientes_trasladados"] or 0)
            for k in RECURSOS:
                s[k] += int(r[k] or 0)
            try:
                hi = r.get("hora_inicio")
                hf = r.get("hora_fin")
                if hi and hf:
                    min_in  = _to_minutes(hi)
                    min_fin = _to_minutes(hf)
                    if min_in is not None and min_fin is not None:
                        diff_min = min_fin - min_in
                        if diff_min < 0:
                            diff_min += 24 * 60   # cruce de medianoche
                        s["horas"] += diff_min / 60.0
            except Exception:
                pass

        for k in stats_servicios:
            stats_servicios[k]["horas"] = round(stats_servicios[k]["horas"], 2)

        # ── Total de recursos globales ──
        total_recursos = {k: sum(s[k] for s in stats_servicios.values()) for k in RECURSOS}
        total_recursos["total_personal"]        = sum(s["total_personal"]        for s in stats_servicios.values())
        total_recursos["pacientes_atendidos"]   = sum(s["pacientes_atendidos"]   for s in stats_servicios.values())
        total_recursos["pacientes_trasladados"] = sum(s["pacientes_trasladados"] for s in stats_servicios.values())
        total_recursos["horas"]                 = round(sum(s["horas"] for s in stats_servicios.values()), 2)

        return render_template("estadisticas_reporte_actividades.html",
                               registros=avalados,
                               stats_servicios=stats_servicios,
                               resumen_estados=resumen_estados,
                               total_recursos=total_recursos)

    @app.route("/formularios/reporte_actividades/exportar_excel")
    @login_required
    def exportar_excel_reporte_actividades():
        es_admin_vol = session["usuario"].get("rol") == "admin"
        formularios_acceso = session["usuario"].get("formularios_acceso", {})
        tiene_permiso = es_admin_vol

        if not es_admin_vol:
            if isinstance(formularios_acceso, dict):
                for perfiles_acc in formularios_acceso.values():
                    if 'voluntariado_admin' in perfiles_acc:
                        es_admin_vol = True
                    if 'reporte_actividades' in perfiles_acc:
                        tiene_permiso = True
            elif isinstance(formularios_acceso, list):
                if 'voluntariado_admin' in formularios_acceso:
                    es_admin_vol = True
                if 'reporte_actividades' in formularios_acceso:
                    tiene_permiso = True

        if not es_admin_vol and not tiene_permiso:
            flash("No tienes permisos para exportar reportes de actividades.", "error")
            return redirect(url_for("dashboard"))

        fecha_desde = request.args.get('fecha_desde', '').strip()
        fecha_hasta = request.args.get('fecha_hasta', '').strip()

        conn = get_db()
        cursor = conn.cursor()

        sql = "SELECT * FROM reporte_actividades"
        params = []
        where_clauses = []
        if fecha_desde:
            where_clauses.append("fecha_evento >= %s")
            params.append(fecha_desde)
        if fecha_hasta:
            where_clauses.append("fecha_evento <= %s")
            params.append(fecha_hasta)

        if where_clauses:
            sql += " WHERE " + " AND ".join(where_clauses)
        sql += " ORDER BY fecha_evento DESC, id DESC"

        cursor.execute(sql, tuple(params))
        registros = cursor.fetchall()
        conn.close()

        if registros and type(registros[0]) is tuple:
            cols = [c[0] for c in cursor.description]
            registros = [dict(zip(cols, r)) for r in registros]
        elif registros and type(registros[0]) is not dict:
            registros = [dict(r) for r in registros]

        def _fmt_time_str(val):
            if not val:
                return "—"
            if isinstance(val, (datetime, time)):
                return val.strftime("%H:%M")
            if isinstance(val, timedelta):
                total_sec = int(val.total_seconds())
                h = (total_sec // 3600) % 24
                m = (total_sec % 3600) // 60
                return f"{h:02d}:{m:02d}"
            s = str(val).strip()
            return s if s else "—"

        def _fmt_date_str(val):
            if not val:
                return "—"
            if isinstance(val, (datetime, date)):
                return val.strftime("%Y-%m-%d")
            s = str(val).strip()
            return s if s else "—"

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Reporte Actividades"
        ws.views.sheetView[0].showGridLines = True

        # Encabezado principal del archivo Excel
        ws.merge_cells('A1:AA1')
        title_cell = ws['A1']
        title_cell.value = "REPORTE DE ACTIVIDADES Y EVENTOS"
        title_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
        title_cell.fill = PatternFill("solid", fgColor="1E3A8A")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 36

        # Subtítulo con rango y metadatos
        ws.merge_cells('A2:AA2')
        info_cell = ws['A2']
        rango_texto = "Histórico Completo"
        if fecha_desde and fecha_hasta:
            rango_texto = f"Desde: {fecha_desde}   Hasta: {fecha_hasta}"
        elif fecha_desde:
            rango_texto = f"Desde: {fecha_desde}"
        elif fecha_hasta:
            rango_texto = f"Hasta: {fecha_hasta}"
        info_cell.value = f"Rango: {rango_texto}   |   Generado: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}   |   Total Registros: {len(registros)}"
        info_cell.font = Font(name="Calibri", size=10, italic=True, color="1E3A8A")
        info_cell.fill = PatternFill("solid", fgColor="EFF6FF")
        info_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 22

        headers = [
            ("ID", 8),
            ("Fecha Evento", 14),
            ("Nombre del Evento", 32),
            ("Lugar / Ubicación", 26),
            ("Tipo de Servicio", 22),
            ("Hora Inicio", 13),
            ("Hora Fin", 13),
            ("Total Personal", 14),
            ("Pacientes Atendidos", 18),
            ("Pacientes Trasladados", 18),
            ("Ambulancia TAB", 15),
            ("Ambulancia TAM", 15),
            ("PASM", 12),
            ("PASB", 12),
            ("Eq. Intervención", 16),
            ("Moto APH", 12),
            ("Unidad Rescate", 15),
            ("Unidad Logística", 16),
            ("Observaciones", 38),
            ("Estado", 14),
            ("Registrado Por", 28),
            ("Documento Registrador", 20),
            ("Perfil Registrador", 20),
            ("Fecha Registro", 20),
            ("Avalado Por", 28),
            ("Documento Avalador", 20),
            ("Fecha Aval", 20)
        ]

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="2563EB")
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1")
        )
        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")

        ws.row_dimensions[4].height = 26
        for col_idx, (header_text, _) in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col_idx, value=header_text)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = align_center
            cell.border = thin_border

        fill_white = PatternFill("solid", fgColor="FFFFFF")
        fill_zebra = PatternFill("solid", fgColor="F8FAFC")

        font_avalado = Font(name="Calibri", size=10, bold=True, color="065F46")
        fill_avalado = PatternFill("solid", fgColor="D1FAE5")
        font_anulado = Font(name="Calibri", size=10, bold=True, color="991B1B")
        fill_anulado = PatternFill("solid", fgColor="FEE2E2")
        font_pendiente = Font(name="Calibri", size=10, bold=True, color="92400E")
        fill_pendiente = PatternFill("solid", fgColor="FEF3C7")

        data_font = Font(name="Calibri", size=10)

        for row_idx, r in enumerate(registros, 5):
            ws.row_dimensions[row_idx].height = 22
            current_fill = fill_zebra if row_idx % 2 == 0 else fill_white

            estado_raw = r.get("estado") or "Pendiente"
            if estado_raw == "Inactivo":
                estado_raw = "Anulado"

            row_values = [
                r.get("id"),
                _fmt_date_str(r.get("fecha_evento")),
                r.get("nombre_evento") or "—",
                r.get("lugar_evento") or "—",
                r.get("tipo_servicio") or "—",
                _fmt_time_str(r.get("hora_inicio")),
                _fmt_time_str(r.get("hora_fin")),
                int(r.get("total_personal") or 0),
                int(r.get("pacientes_atendidos") or 0),
                int(r.get("pacientes_trasladados") or 0),
                int(r.get("ambulancia_tab") or 0),
                int(r.get("ambulancia_tam") or 0),
                int(r.get("pasm") or 0),
                int(r.get("pasb") or 0),
                int(r.get("equipos_intervencion") or 0),
                int(r.get("moto_aph") or 0),
                int(r.get("unidad_rescate") or 0),
                int(r.get("unidad_logistica") or 0),
                r.get("observaciones") or "—",
                estado_raw,
                r.get("registrado_por") or "—",
                r.get("registrado_por_identificacion") or "—",
                r.get("perfil_registrador") or "—",
                str(r.get("fecha_registro") or "—"),
                r.get("avalado_por") or "—",
                r.get("avalado_por_identificacion") or "—",
                str(r.get("fecha_aval") or "—")
            ]

            for col_idx, val in enumerate(row_values, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val if val is not None else "—")
                cell.font = data_font
                cell.border = thin_border
                cell.fill = current_fill

                # Centrar columnas numéricas, fechas, horas y estados; alinear texto descriptivo a la izquierda
                if col_idx in (1, 2, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17, 18, 20, 22, 23, 24, 26, 27):
                    cell.alignment = align_center
                else:
                    cell.alignment = align_left

                # Estilo tipo insignia para el estado
                if col_idx == 20:
                    if estado_raw == "Avalado":
                        cell.font = font_avalado
                        cell.fill = fill_avalado
                    elif estado_raw in ("Anulado", "Inactivo"):
                        cell.font = font_anulado
                        cell.fill = fill_anulado
                    else:
                        cell.font = font_pendiente
                        cell.fill = fill_pendiente

        for col_idx, (_, width) in enumerate(headers, 1):
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = width

        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        fecha_archivo = datetime.now().strftime("%Y%m%d")
        return send_file(
            buf,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=f"Reportes_Actividades_{fecha_archivo}.xlsx"
        )


