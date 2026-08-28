from flask import render_template, request, session, redirect, url_for, flash
from utils import login_required
from db import get_db

def register_routes(app):
    @app.route("/programacion_operativa", methods=["GET", "POST"])
    @login_required
    def programacion_operativa():
        if not session.get("usuario") or (not session["usuario"].get("permiso_programacion_operativa") and str(session["usuario"].get("id")) != '1'):
            flash("No tienes permiso para acceder a Programación Operativa.", "error")
            return redirect(url_for("dashboard"))
            
        conn = get_db()
        if request.method == "POST":
            tipo_evento = request.form.get("tipo_evento")
            nombre_evento = request.form.get("nombre_evento", "").strip()
            fecha = request.form.get("fecha")
            hora_inicio = request.form.get("hora_inicio", "").strip()
            hora_finalizacion = request.form.get("hora_finalizacion", "").strip()
            lugar = request.form.get("lugar", "").strip()
            contacto = request.form.get("contacto", "").strip()
            coordina = request.form.get("coordina", "").strip()
            
            if tipo_evento in ["FUTBOL", "Concierto Estadio", "Concierto Macarena"]:
                contacto = request.form.get("representante_pcu", "").strip()
                coordina = request.form.get("coordinador_terreno", "").strip()
                
            recursos_tecnicos = request.form.get("recursos_tecnicos", "").strip()
            registrado_por = session["usuario"]["id"]
            
            try:
                conn.execute("ALTER TABLE programacion_operativa ADD COLUMN estado VARCHAR(50) DEFAULT 'FINALIZADO'")
            except:
                pass
                
            estado = request.form.get("estado_programacion", "FINALIZADO")
            
            columnas_layout = request.form.get("columnas_layout", "3")
            try:
                columnas_layout = int(columnas_layout)
            except:
                columnas_layout = 3
                
            hora_inicio_pcu = request.form.get("hora_inicio_pcu", "").strip()
            hora_apertura_puertas = request.form.get("hora_apertura_puertas", "").strip()
            hora_inicio_evento = request.form.get("hora_inicio_evento", "").strip()
            hora_finalizacion_pcu = request.form.get("hora_finalizacion_pcu", "").strip()
            hora_llegada_aph = request.form.get("hora_llegada_aph", "").strip()
            hora_retiro_aph = request.form.get("hora_retiro_aph", "").strip()
            cantidad_recurso_humano = request.form.get("cantidad_recurso_humano", "").strip()
            cantidad_ambulancias = request.form.get("cantidad_ambulancias", "").strip()
            total_asistentes = request.form.get("total_asistentes", "").strip()
            total_pacientes = request.form.get("total_pacientes", "").strip()
            programacion_id = request.form.get("programacion_id")
            
            if programacion_id:
                # Update existing
                conn.execute("""
                    UPDATE programacion_operativa SET
                    tipo_evento=?, nombre_evento=?, fecha=?, hora_inicio=?, hora_finalizacion=?, lugar=?, contacto=?, coordina=?, recursos_tecnicos=?, registrado_por=?, columnas_layout=?, estado=?,
                    hora_inicio_pcu=?, hora_apertura_puertas=?, hora_inicio_evento=?, hora_finalizacion_pcu=?, hora_llegada_aph=?, hora_retiro_aph=?, cantidad_recurso_humano=?, cantidad_ambulancias=?, total_asistentes=?, total_pacientes=?
                    WHERE id=?
                """, (tipo_evento, nombre_evento, fecha, hora_inicio, hora_finalizacion, lugar, contacto, coordina, recursos_tecnicos, registrado_por, columnas_layout, estado,
                      hora_inicio_pcu, hora_apertura_puertas, hora_inicio_evento, hora_finalizacion_pcu, hora_llegada_aph, hora_retiro_aph, cantidad_recurso_humano, cantidad_ambulancias, total_asistentes, total_pacientes, programacion_id))
                
                # Delete existing members to replace them
                conn.execute("DELETE FROM programacion_operativa_integrantes WHERE programacion_id=?", (programacion_id,))
            else:
                cursor = conn.execute("""
                    INSERT INTO programacion_operativa 
                    (tipo_evento, nombre_evento, fecha, hora_inicio, hora_finalizacion, lugar, contacto, coordina, recursos_tecnicos, registrado_por, columnas_layout, estado,
                     hora_inicio_pcu, hora_apertura_puertas, hora_inicio_evento, hora_finalizacion_pcu, hora_llegada_aph, hora_retiro_aph, cantidad_recurso_humano, cantidad_ambulancias, total_asistentes, total_pacientes)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (tipo_evento, nombre_evento, fecha, hora_inicio, hora_finalizacion, lugar, contacto, coordina, recursos_tecnicos, registrado_por, columnas_layout, estado,
                      hora_inicio_pcu, hora_apertura_puertas, hora_inicio_evento, hora_finalizacion_pcu, hora_llegada_aph, hora_retiro_aph, cantidad_recurso_humano, cantidad_ambulancias, total_asistentes, total_pacientes))
                programacion_id = cursor.lastrowid
            
            
            if tipo_evento == "Tripulacion Basica":
                nombres = request.form.getlist("integrante_nombre[]")
                roles = request.form.getlist("integrante_rol[]")
                
                basica_tipo = request.form.get("basica_tipo_ambulancia", "").strip()
                basica_num = request.form.get("basica_numero_movil", "").strip()
                
                u_nombre = ""
                if basica_tipo or basica_num:
                    if basica_tipo:
                        u_nombre = f"AMB. {basica_tipo}"
                        if basica_num:
                            u_nombre += f" {basica_num}"
                    else:
                        u_nombre = basica_num
                
                for i, nombre in enumerate(nombres):
                    if nombre.strip():
                        rol = roles[i] if i < len(roles) else ""
                        conn.execute("""
                            INSERT INTO programacion_operativa_integrantes (programacion_id, nombre, rol_variable, orden, unidad_nombre)
                            VALUES (?, ?, ?, ?, ?)
                        """, (programacion_id, nombre.strip(), rol.strip(), i+1, u_nombre))
            elif tipo_evento in ["Operativo Completo", "FUTBOL", "Concierto Estadio", "Concierto Macarena"]:
                unidad_nombres = request.form.getlist("unidad_nombre[]")
                unidad_tipos = request.form.getlist("unidad_tipo[]")
                unidad_counts = request.form.getlist("unidad_campos_count[]")
                u_roles = request.form.getlist("unidad_integrante_rol[]")
                u_nombres = request.form.getlist("unidad_integrante_nombre[]")
                
                offset = 0
                for u_idx, u_count_str in enumerate(unidad_counts):
                    try:
                        u_count = int(u_count_str)
                    except:
                        u_count = 0
                        
                    u_nombre = unidad_nombres[u_idx] if u_idx < len(unidad_nombres) else ""
                    u_tipo = unidad_tipos[u_idx] if u_idx < len(unidad_tipos) else ""
                    
                    for i in range(u_count):
                        idx = offset + i
                        if idx < len(u_nombres):
                            nombre = u_nombres[idx].strip()
                            rol = u_roles[idx].strip() if idx < len(u_roles) else ""
                            if nombre:
                                conn.execute("""
                                    INSERT INTO programacion_operativa_integrantes 
                                    (programacion_id, nombre, rol_variable, orden, unidad_nombre)
                                    VALUES (?, ?, ?, ?, ?)
                                """, (programacion_id, nombre, rol, i+1, u_nombre))
                    offset += u_count
            conn.commit()
            conn.close()
            flash("Programación operativa guardada correctamente.", "success")
            return redirect(url_for("programacion_operativa_explorador"))
            
        # Obtener personal de la nueva tabla personal_operativo
        personal_op = conn.execute("SELECT nombres, apellidos, perfiles FROM personal_operativo ORDER BY nombres, apellidos").fetchall()
        empleados_por_perfil = {
            "COND": [],
            "MED": [],
            "APH": [],
            "AUX ENF": [],
            "ENF": [],
            "SOC": [],
            "SIN_ROL": []
        }
        for p in personal_op:
            nombre = f"{p['nombres']} {p['apellidos']}".strip()
            perfiles_str = p['perfiles'] or ""
            perfiles_list = [perf.strip() for perf in perfiles_str.split(',') if perf.strip()]
            
            if not perfiles_list:
                empleados_por_perfil["SIN_ROL"].append(nombre)
            else:
                for perf in perfiles_list:
                    if perf in empleados_por_perfil:
                        empleados_por_perfil[perf].append(nombre)
                    else:
                        if nombre not in empleados_por_perfil["SIN_ROL"]:
                            empleados_por_perfil["SIN_ROL"].append(nombre)
                            
        # Asegurar valores únicos y ordenados por si acaso
        for k in empleados_por_perfil:
            empleados_por_perfil[k] = sorted(list(set(empleados_por_perfil[k])))
                
        conn.close()
        return render_template("programacion_operativa.html", usuario=session["usuario"], empleados_por_perfil=empleados_por_perfil)

    @app.route("/programacion_operativa/editar/<int:id>")
    @login_required
    def programacion_operativa_editar(id):
        if not session.get("usuario") or (not session["usuario"].get("permiso_programacion_operativa") and str(session["usuario"].get("id")) != '1'):
            flash("No tienes permiso.", "error")
            return redirect(url_for("dashboard"))
            
        conn = get_db()
        evento = conn.execute("SELECT * FROM programacion_operativa WHERE id = ?", (id,)).fetchone()
        if not evento:
            conn.close()
            flash("Evento no encontrado.", "error")
            return redirect(url_for("programacion_operativa_explorador"))
            
        integrantes = conn.execute("SELECT * FROM programacion_operativa_integrantes WHERE programacion_id = ? ORDER BY id", (id,)).fetchall()
        
        import json
        evento_dict = dict(evento)
        integrantes_list = [dict(i) for i in integrantes]
        
        personal_op = conn.execute("SELECT nombres, apellidos, perfiles FROM personal_operativo ORDER BY nombres, apellidos").fetchall()
        empleados_por_perfil = {
            "COND": [], "MED": [], "APH": [], "AUX ENF": [], "ENF": [], "SOC": [], "SIN_ROL": []
        }
        for p in personal_op:
            nombre = f"{p['nombres']} {p['apellidos']}".strip()
            perfiles_str = p['perfiles'] or ""
            perfiles_list = [perf.strip() for perf in perfiles_str.split(',') if perf.strip()]
            
            if not perfiles_list:
                empleados_por_perfil["SIN_ROL"].append(nombre)
            else:
                for perf in perfiles_list:
                    if perf in empleados_por_perfil:
                        empleados_por_perfil[perf].append(nombre)
                    else:
                        if nombre not in empleados_por_perfil["SIN_ROL"]:
                            empleados_por_perfil["SIN_ROL"].append(nombre)
                            
        for k in empleados_por_perfil:
            empleados_por_perfil[k] = sorted(list(set(empleados_por_perfil[k])))
            
        conn.close()
        
        return render_template(
            "programacion_operativa.html", 
            usuario=session["usuario"], 
            empleados_por_perfil=empleados_por_perfil,
            evento_a_editar=evento_dict,
            integrantes_a_editar=json.dumps(integrantes_list)
        )


    @app.route("/api/verificar_disponibilidad", methods=["POST"])
    @login_required
    def verificar_disponibilidad():
        from flask import jsonify
        data = request.get_json()
        nombre = data.get("nombre", "").strip()
        fecha = data.get("fecha", "").strip()
        hora_inicio_nueva = data.get("hora_inicio", "").strip()
        hora_fin_nueva = data.get("hora_finalizacion", "").strip()
        
        if not nombre or not fecha or not hora_inicio_nueva:
            return jsonify({"cruce": False})
            
        conn = get_db()
        # Verificar si hay programaciones para ese nombre en esa fecha
        query = """
            SELECT po.nombre_evento, po.hora_inicio, po.hora_finalizacion
            FROM programacion_operativa po
            JOIN programacion_operativa_integrantes poi ON po.id = poi.programacion_id
            WHERE poi.nombre = ? AND po.fecha = ?
        """
        rows = conn.execute(query, (nombre, fecha)).fetchall()
        conn.close()
        
        eventos_cruce = []
        for r in rows:
            h_ini_exist = r['hora_inicio']
            h_fin_exist = r['hora_finalizacion']
            
            # Simple lógica para cruce
            fin_exist_cmp = h_fin_exist if h_fin_exist else "24:00"
            if fin_exist_cmp == "00:00": fin_exist_cmp = "24:00"
            
            fin_nuevo_cmp = hora_fin_nueva if hora_fin_nueva else "24:00"
            if fin_nuevo_cmp == "00:00": fin_nuevo_cmp = "24:00"
            
            if h_ini_exist < fin_nuevo_cmp and fin_exist_cmp > hora_inicio_nueva:
                horario = f"de {h_ini_exist}"
                if h_fin_exist:
                    horario += f" a {h_fin_exist}"
                eventos_cruce.append(f"'{r['nombre_evento']}' ({horario})")
            
        if eventos_cruce:
            msj = f"El integrante {nombre} tiene un cruce de horario en: " + ", ".join(eventos_cruce) + ". ¿Deseas mantenerlo asignado?"
            return jsonify({"cruce": True, "mensaje": msj})
            
        return jsonify({"cruce": False})

    @app.route("/programacion_operativa/explorador")
    @login_required
    def programacion_operativa_explorador():
        if not session.get("usuario") or (not session["usuario"].get("permiso_programacion_operativa") and str(session["usuario"].get("id")) != '1'):
            flash("No tienes permiso.", "error")
            return redirect(url_for("dashboard"))
            
        conn = get_db()
        fecha_desde = request.args.get('fecha_desde', '').strip()
        fecha_hasta = request.args.get('fecha_hasta', '').strip()
        
        if fecha_desde and fecha_hasta:
            eventos = conn.execute("SELECT * FROM programacion_operativa WHERE fecha >= ? AND fecha <= ? ORDER BY fecha DESC, id DESC", (fecha_desde, fecha_hasta)).fetchall()
        elif fecha_desde:
            eventos = conn.execute("SELECT * FROM programacion_operativa WHERE fecha >= ? ORDER BY fecha DESC, id DESC", (fecha_desde,)).fetchall()
        elif fecha_hasta:
            eventos = conn.execute("SELECT * FROM programacion_operativa WHERE fecha <= ? ORDER BY fecha DESC, id DESC", (fecha_hasta,)).fetchall()
        else:
            eventos = conn.execute("SELECT * FROM programacion_operativa ORDER BY id DESC").fetchall()
        
        for evento in eventos:
            evento["integrantes"] = conn.execute("SELECT * FROM programacion_operativa_integrantes WHERE programacion_id = ? ORDER BY orden", (evento["id"],)).fetchall()
            
        conn.close()
        return render_template("programacion_operativa_explorador.html", eventos=eventos, usuario=session["usuario"], fecha_desde=fecha_desde, fecha_hasta=fecha_hasta)

    @app.route("/programacion_operativa/imprimir/<int:evento_id>")
    @login_required
    def programacion_operativa_imprimir(evento_id):
        if not session.get("usuario") or (not session["usuario"].get("permiso_programacion_operativa") and str(session["usuario"].get("id")) != '1'):
            return "No tienes permiso.", 403
            
        conn = get_db()
        evento = conn.execute("SELECT * FROM programacion_operativa WHERE id = ?", (evento_id,)).fetchone()
        if not evento:
            conn.close()
            return "Evento no encontrado", 404
            
        integrantes = conn.execute("SELECT * FROM programacion_operativa_integrantes WHERE programacion_id = ? ORDER BY orden", (evento["id"],)).fetchall()
        
        # Enriquecer integrantes con cédula y registro para formato FUTBOL/Concierto/Excel
        personal_db = conn.execute("SELECT nombres, apellidos, cedula, codigo_fecha, registro, perfiles FROM personal_operativo").fetchall()
        conn.close()
        
        personal_lookup = {}
        for p in personal_db:
            nombre_completo = f"{p['nombres']} {p['apellidos']}".strip().upper()
            personal_lookup[nombre_completo] = {
                'cedula': p['cedula'],
                'codigo_fecha': p['codigo_fecha'],
                'registro': p['registro'],
                'perfiles': p['perfiles']
            }
            nombre_solo = p['nombres'].strip().upper()
            if nombre_solo and nombre_solo not in personal_lookup:
                personal_lookup[nombre_solo] = personal_lookup[nombre_completo]

        # Convertir a lista de diccionarios para poder modificarlos
        integrantes = [dict(ing) for ing in integrantes]
        for ing in integrantes:
            nombre_ing = ing.get('nombre', '').strip().upper()
            info_personal = personal_lookup.get(nombre_ing, {})
            ing['cedula'] = info_personal.get('cedula', '')
            ing['registro'] = info_personal.get('registro', '')
            ing['perfiles'] = info_personal.get('perfiles', '')
        
        # Formatear la fecha
        import locale
        from datetime import datetime
        try:
            locale.setlocale(locale.LC_TIME, 'es_ES.UTF-8')
        except:
            try:
                locale.setlocale(locale.LC_TIME, 'es_CO.utf8')
            except:
                pass # Fallback to default if locales not found
                
        fecha_obj = evento['fecha']
        if isinstance(fecha_obj, str):
            try:
                fecha_obj = datetime.strptime(fecha_obj, '%Y-%m-%d')
            except:
                pass
                
        fecha_formateada = ""
        mes_ano = ""
        if isinstance(fecha_obj, datetime) or hasattr(fecha_obj, 'strftime'):
            # Convert to uppercase
            meses = ['ENERO', 'FEBRERO', 'MARZO', 'ABRIL', 'MAYO', 'JUNIO', 'JULIO', 'AGOSTO', 'SEPTIEMBRE', 'OCTUBRE', 'NOVIEMBRE', 'DICIEMBRE']
            dias = ['LUNES', 'MARTES', 'MIÉRCOLES', 'JUEVES', 'VIERNES', 'SÁBADO', 'DOMINGO']
            mes_str = meses[fecha_obj.month - 1]
            dia_semana_str = dias[fecha_obj.weekday()]
            
            fecha_formateada = f"{mes_str} {fecha_obj.day} ({dia_semana_str})"
            mes_ano = f"{mes_str} {fecha_obj.year}"
            
        def format_12h(time_str):
            if not time_str or time_str.strip() == '':
                return ''
            try:
                # Si ya tiene am/pm no hacer nada
                if 'am' in time_str.lower() or 'pm' in time_str.lower():
                    return time_str
                parts = time_str.split(':')
                h = int(parts[0])
                m = int(parts[1])
                suffix = 'am' if h < 12 else 'pm'
                h_12 = h if h <= 12 else h - 12
                if h_12 == 0: h_12 = 12
                return f"{h_12}:{m:02d} {suffix}"
            except:
                return time_str
                
        # Modificamos el diccionario de evento para mostrar las horas formateadas en la plantilla
        evento_dict = dict(evento)
        evento_dict['hora_inicio'] = format_12h(evento_dict.get('hora_inicio', ''))
        evento_dict['hora_finalizacion'] = format_12h(evento_dict.get('hora_finalizacion', ''))

        unidades = []
        if evento_dict.get('tipo_evento') in ['Operativo Completo', 'FUTBOL', 'Concierto Estadio', 'Concierto Macarena']:
            last_u_nombre = None
            current_unidad = None
            for ing in integrantes:
                u_nombre = ing.get('unidad_nombre') or 'UNIDAD'
                
                # Para evitar agrupar unidades del mismo nombre que están separadas, y agrupar las contiguas
                if u_nombre != last_u_nombre:
                    if current_unidad:
                        unidades.append(current_unidad)
                    current_unidad = {'nombre': u_nombre, 'integrantes': []}
                    last_u_nombre = u_nombre
                    
                if ing.get('nombre') != 'PAGE_BREAK_DUMMY':
                    current_unidad['integrantes'].append(ing)
            
            if current_unidad:
                unidades.append(current_unidad)

        # Alternative date format: DOMINGO 24 DE MAYO
        fecha_completo = ""
        if isinstance(fecha_obj, datetime) or hasattr(fecha_obj, 'strftime'):
            fecha_completo = f"{dia_semana_str} {fecha_obj.day} DE {mes_str}"

        context = {
            'evento': evento_dict,
            'integrantes': integrantes,
            'unidades': unidades,
            'fecha_formateada': fecha_formateada,
            'fecha_completo': fecha_completo,
            'mes_ano': mes_ano
        }
        
        return render_template("programacion_operativa_imprimir.html", **context)

    @app.route("/programacion_operativa/confirmar", methods=["POST"])
    @login_required
    def programacion_operativa_confirmar():
        if not session.get("usuario") or (not session["usuario"].get("permiso_programacion_operativa") and str(session["usuario"].get("id")) != '1'):
            return "No autorizado", 403
            
        evento_id = request.form.get("evento_id")
        asistentes_ids = request.form.getlist("asistio[]")
        
        hora_inicio = request.form.get("hora_inicio", "").strip()
        hora_finalizacion = request.form.get("hora_finalizacion", "").strip()
        
        try:
            val_asis = request.form.get("tarifa_asistencial", "0").replace(".", "").replace(",", "")
            tarifa_asistencial = int(val_asis) if val_asis else 0
        except:
            tarifa_asistencial = 0
            
        try:
            val_cond = request.form.get("tarifa_conductor", "0").replace(".", "").replace(",", "")
            tarifa_conductor = int(val_cond) if val_cond else 0
        except:
            tarifa_conductor = 0
        
        if not evento_id:
            flash("Error: No se proporcionó el ID del evento.", "error")
            return redirect(url_for("programacion_operativa_explorador"))
            
        conn = get_db()
        try:
            # Marcar el evento como confirmado y guardar horas/tarifas
            conn.execute("""
                UPDATE programacion_operativa 
                SET confirmado = 1, hora_inicio = ?, hora_finalizacion = ?, tarifa_asistencial = ?, tarifa_conductor = ? 
                WHERE id = ?
            """, (hora_inicio, hora_finalizacion, tarifa_asistencial, tarifa_conductor, evento_id))
            
            # Todos los integrantes de este evento se marcan como no asistieron por defecto
            conn.execute("UPDATE programacion_operativa_integrantes SET asistio = 0 WHERE programacion_id = ?", (evento_id,))
            
            # Y los que vinieron en la lista (los que dejaron el checkbox marcado) se marcan como asistieron
            if asistentes_ids:
                placeholders = ",".join("?" * len(asistentes_ids))
                conn.execute(f"UPDATE programacion_operativa_integrantes SET asistio = 1 WHERE programacion_id = ? AND id IN ({placeholders})", [evento_id] + asistentes_ids)
                
            conn.commit()
            flash("Asistencia confirmada correctamente.", "success")
        except Exception as e:
            conn.rollback()
            flash(f"Error al confirmar asistencia: {e}", "error")
        finally:
            conn.close()
            
        return redirect(url_for("programacion_operativa_explorador"))

    from flask import jsonify

    @app.route("/api/check_programacion_conflict", methods=["POST"])
    @login_required
    def check_programacion_conflict():
        data = request.get_json()
        if not data:
            return jsonify({"conflictos": []})
            
        fecha = data.get('fecha')
        hora_inicio = data.get('hora_inicio')
        hora_finalizacion = data.get('hora_finalizacion')
        nombres = data.get('nombres', [])
        ignore_evento_id = data.get('ignore_evento_id')
        
        if not fecha or not nombres:
            return jsonify({"conflictos": []})
            
        conn = get_db()
        try:
            query = """
                SELECT po.id, po.nombre_evento, po.hora_inicio, po.hora_finalizacion, poi.nombre 
                FROM programacion_operativa po
                JOIN programacion_operativa_integrantes poi ON po.id = poi.programacion_id
                WHERE po.fecha = ?
            """
            params = [fecha]
            
            if ignore_evento_id:
                query += " AND po.id != ?"
                params.append(ignore_evento_id)
                
            rows = conn.execute(query, params).fetchall()
            
            conflictos = []
            for r in rows:
                if r['nombre'] in nombres:
                    h_i = r['hora_inicio']
                    h_f = r['hora_finalizacion'] or "23:59"
                    
                    req_h_i = hora_inicio or "00:00"
                    req_h_f = hora_finalizacion or "23:59"
                    
                    if req_h_i < h_f and req_h_f > h_i:
                        conflictos.append({
                            "nombre": r['nombre'],
                            "evento": r['nombre_evento'],
                            "hora_inicio": h_i,
                            "hora_finalizacion": h_f
                        })
            return jsonify({"conflictos": conflictos})
        except Exception as e:
            print("Error checking conflict:", e)
            return jsonify({"conflictos": []})
        finally:
            conn.close()

    @app.route("/programacion_operativa/editar_integrantes", methods=["POST"])
    @login_required
    def programacion_operativa_editar_integrantes():
        if not session.get("usuario") or (not session["usuario"].get("permiso_programacion_operativa") and str(session["usuario"].get("id")) != '1'):
            return "No autorizado", 403
            
        evento_id = request.form.get("evento_id")
        if not evento_id:
            flash("Error: No se proporcionó el ID del evento.", "error")
            return redirect(url_for("programacion_operativa_explorador"))
            
        integrante_ids = request.form.getlist("integrante_id[]")
        nombres = request.form.getlist("nombre[]")
        roles = request.form.getlist("rol_variable[]")
        unidades = request.form.getlist("unidad_nombre[]")
        
        conn = get_db()
        try:
            # Delete those not in the current list
            valid_ids = [i for i in integrante_ids if i.isdigit() and i != '0']
            if valid_ids:
                placeholders = ",".join("?" * len(valid_ids))
                conn.execute(f"DELETE FROM programacion_operativa_integrantes WHERE programacion_id = ? AND id NOT IN ({placeholders})", [evento_id] + valid_ids)
            else:
                conn.execute("DELETE FROM programacion_operativa_integrantes WHERE programacion_id = ?", (evento_id,))
                
            # Update or Insert
            for i in range(len(nombres)):
                nombre = nombres[i].strip()
                if not nombre:
                    continue
                rol = roles[i].strip() if i < len(roles) else ""
                unidad = unidades[i].strip() if i < len(unidades) else ""
                i_id = integrante_ids[i] if i < len(integrante_ids) else "0"
                
                if i_id.isdigit() and i_id != '0':
                    # Update
                    conn.execute("""
                        UPDATE programacion_operativa_integrantes 
                        SET nombre = ?, rol_variable = ?, unidad_nombre = ?
                        WHERE id = ? AND programacion_id = ?
                    """, (nombre, rol, unidad, i_id, evento_id))
                else:
                    # Insert
                    conn.execute("""
                        INSERT INTO programacion_operativa_integrantes (programacion_id, nombre, rol_variable, unidad_nombre, orden)
                        VALUES (?, ?, ?, ?, ?)
                    """, (evento_id, nombre, rol, unidad, i+1))
                    
            conn.commit()
            flash("Integrantes actualizados correctamente.", "success")
        except Exception as e:
            conn.rollback()
            flash(f"Error al actualizar integrantes: {e}", "error")
        finally:
            conn.close()
            
        return redirect(url_for("programacion_operativa_explorador"))

    @app.route("/programacion_operativa/exportar_pagos")
    @login_required
    def programacion_operativa_exportar_pagos():
        fecha_inicio = request.args.get("fecha_inicio")
        fecha_fin = request.args.get("fecha_fin")
        
        if not fecha_inicio or not fecha_fin:
            flash("Debe seleccionar un rango de fechas.", "error")
            return redirect(url_for("programacion_operativa_explorador"))
            
        conn = get_db()
        query = """
            SELECT 
                po.id as evento_id,
                po.fecha,
                po.nombre_evento,
                po.tarifa_asistencial,
                po.tarifa_conductor,
                poi.nombre as integrante_nombre,
                poi.rol_variable as rol
            FROM programacion_operativa po
            JOIN programacion_operativa_integrantes poi ON po.id = poi.programacion_id
            WHERE po.confirmado = 1 
              AND poi.asistio = 1
              AND po.fecha >= ? 
              AND po.fecha <= ?
            ORDER BY po.fecha ASC, po.id ASC
        """
        filas = conn.execute(query, (fecha_inicio, fecha_fin)).fetchall()
        
        # Obtener datos del último archivo de nómina para mapear nombres a identificacion, nombres, apellidos
        ultima_nomina = conn.execute("SELECT id FROM nomina ORDER BY id DESC LIMIT 1").fetchone()
        mapa_empleados = {}
        if ultima_nomina:
            empleados = conn.execute("SELECT identificacion, nombres, apellidos, codigo FROM nomina_empleados WHERE nomina_id = ?", (ultima_nomina["id"],)).fetchall()
            for emp in empleados:
                # Usar el nombre completo como llave de busqueda simple (asumiendo que poi.nombre coincide con nombres + apellidos o solo nombres)
                llave1 = f"{emp['nombres']} {emp['apellidos']}".strip().upper()
                llave2 = emp['nombres'].strip().upper() if emp['nombres'] else ""
                mapa_empleados[llave1] = emp
                if llave2:
                    mapa_empleados[llave2] = emp

        # Agrupar datos por integrante y evento
        datos_agrupados = {}
        # Pre-llenar con todos los empleados de la nómina
        if ultima_nomina:
            for emp in empleados:
                llave_oficial = f"{emp['nombres']} {emp['apellidos']}".strip().upper()
                datos_agrupados[llave_oficial] = {}

        eventos_unicos = []
        eventos_set = set()
        
        for fila in filas:
            integrante_orig = fila['integrante_nombre'].strip()
            integrante = integrante_orig.upper()
            
            # Mapear al nombre oficial si es posible para evitar duplicados
            if integrante not in datos_agrupados:
                emp_data = mapa_empleados.get(integrante)
                if emp_data:
                    integrante = f"{emp_data['nombres']} {emp_data['apellidos']}".strip().upper()
            
            fecha = fila['fecha']
            nombre_evento = fila['nombre_evento'] or "Evento"
            evento_id = fila['evento_id']
            rol = fila['rol'] or ""
            
            # Llave unica para la columna del evento
            llave_evento = f"{fecha} - {nombre_evento} (#{evento_id})"
            
            if llave_evento not in eventos_set:
                eventos_set.add(llave_evento)
                eventos_unicos.append(llave_evento)
            
            tarifa = 0
            rol_upper = rol.upper().strip()
            if rol_upper == 'MED':
                tarifa = 0
            elif rol_upper in ['COND', 'SOC'] or rol_upper == '':
                tarifa = fila['tarifa_conductor'] or 0
            else:
                tarifa = fila['tarifa_asistencial'] or 0
                
            if integrante not in datos_agrupados:
                datos_agrupados[integrante] = {}
                
            datos_agrupados[integrante][llave_evento] = datos_agrupados[integrante].get(llave_evento, 0) + tarifa

        conn.close()
        
        import io
        import openpyxl
        from flask import send_file
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Pagos"
        
        headers = ['NUMERO DE IDENTIFICACION', 'NOMBRES', 'APELLIDOS', 'CODIGO'] + eventos_unicos + ['TOTAL']
        ws.append(headers)
        
        for integrante, pagos_por_evento in datos_agrupados.items():
            llave_busqueda = integrante.upper()
            emp_data = mapa_empleados.get(llave_busqueda)
            
            if emp_data:
                codigo = emp_data['codigo']
                identificacion = emp_data['identificacion']
                nombres = emp_data['nombres']
                apellidos = emp_data['apellidos']
            else:
                codigo = ""
                identificacion = ""
                # Intentar separar el nombre si no se encuentra
                partes = integrante.split(" ", 1)
                nombres = partes[0]
                apellidos = partes[1] if len(partes) > 1 else ""
                
            fila_excel = [identificacion, nombres, apellidos, codigo]
            
            total = 0
            for evento in eventos_unicos:
                valor = pagos_por_evento.get(evento, 0)
                fila_excel.append(valor)
                total += valor
                
            fila_excel.append(total)
            ws.append(fila_excel)
            
        excel_io = io.BytesIO()
        wb.save(excel_io)
        excel_io.seek(0)
        
        return send_file(
            excel_io,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=f'Exportar_Pagos_{fecha_inicio}_al_{fecha_fin}.xlsx'
        )
